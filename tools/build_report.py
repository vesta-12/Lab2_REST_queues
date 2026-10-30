"""Build the PDF and matching editable DOCX from unchanged evidence images.
Optional authoring dependencies: requirements-report.txt. Run from any directory.
"""
import argparse
import json
from pathlib import Path
import re
from xml.sax.saxutils import escape

from PIL import Image as PILImage
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, PageBreak, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--docx', type=Path, default=ROOT.parent / 'Lab2_Report_Serbina_Sofiya.docx')
    parser.add_argument('--font-dir', type=Path, default=Path('/System/Library/Fonts/Supplemental'))
    args = parser.parse_args()
    for face, filename in [('TNR','Times New Roman.ttf'), ('TNR-Bold','Times New Roman Bold.ttf'), ('TNR-Italic','Times New Roman Italic.ttf')]:
        pdfmetrics.registerFont(TTFont(face, str(args.font_dir / filename)))
    metrics = json.loads((ROOT / 'evidence/metrics.json').read_text())
    summary = re.search(r'(\d+ passed in [\d.]+s)', (ROOT / 'evidence/pytest.txt').read_text()).group(1)
    blocks = [
        ('header', 'Programme: Computing and Digital Technologies'),
        ('header', 'Student: Sofiya Serbina'),
        ('title', 'Laboratory Work No. 2'),
        ('header', 'REST API and Message Queues'),
        ('heading', 'Section A Implementation and Results'),
        ('body', 'The revised Flask API provides synchronous /time and /echo requests and asynchronous text analysis. POST /jobs durably queues a message and returns HTTP 202 with a UUID. An independent worker reads the file queue and stores the result for GET /jobs/<job_id>.'),
        ('image', '01_rest_api_200_400_404.png'),
        ('caption', 'Figure 1. Terminal responses: /time and /echo return 200; invalid input returns 400; an unknown job returns 404.'),
        ('image', '02_job_queued.png'),
        ('caption', 'Figure 2. Job 5d445131-c6c0-4d87-9cc8-ec022e2921a1 is queued with payload "Hello queue".'),
        ('image', '03_api_logs.png'),
        ('caption', 'Figure 3. Flask records POST /jobs with status 202. This log is server evidence; it does not display the client response body.'),
        ('break', ''),
        ('image', '04_job_completed.png'),
        ('caption', 'Figure 4. The same UUID is done: uppercase HELLO QUEUE, character_count 11, word_count 2, with lifecycle timestamps.'),
        ('heading', 'Section B Asynchronous Processing and Worker Recovery'),
        ('body', 'HTTP 202 confirms acceptance, not completion. Jobs remain queued while the worker is stopped. The lifecycle is queued, processing, then done or failed. The worker claims a file, computes text statistics, saves the terminal result and only then removes the active queue record (ACK). Startup recovers incomplete claims; an existing saved result takes precedence over a leftover claim.'),
        ('body', 'FIFO uses creation time and UUID for ties. Recovery provides at-least-once execution. Repeated text analysis is safe because it is deterministic and has no external side effects. A crash before ACK must not lose an already saved result.'),
        ('heading', 'Testing and Verification'),
        ('image', '05_pytest_41_passed.png'),
        ('caption', 'Figure 5. Student terminal run: 41 passed in 30.36s. The separate final rerun is recorded in evidence/pytest.txt.'),
        ('body', f'The final rerun reported {summary}, with 0 failed. Ten jobs persisted without a worker for {metrics["worker_absent_seconds"]:.3f} seconds; all ten resumed with correct results in {metrics["resumed_seconds"]:.4f} seconds. The p95 of 100 test-client submissions was {metrics["api_p95_seconds"]*1000:.3f} ms. Tests also verified recovery after a killed process and preservation of the saved result after a failure before ACK.'),
        ('heading', 'Conclusion and Limitations'),
        ('body', 'The tested implementation supports asynchronous processing and recovery on one local POSIX host with one worker. It does not provide authentication, distributed storage or a guarantee against every hardware failure. Original code and specification shared commit c5837d7. The honest remediation sequence was specification c4772f7, implementation 8652cab, then verification 7be01cd. OpenAI Codex assisted the remediation; the student reviews and explains the code.'),
    ]
    styles = {
        'body': ParagraphStyle('body',fontName='TNR',fontSize=14,leading=16,spaceAfter=5),
        'header': ParagraphStyle('header',fontName='TNR',fontSize=12,leading=14,alignment=1,spaceAfter=2),
        'title': ParagraphStyle('title',fontName='TNR-Bold',fontSize=14,leading=16,alignment=1,spaceAfter=2),
        'heading': ParagraphStyle('heading',fontName='TNR-Bold',fontSize=14,leading=16,spaceBefore=5,spaceAfter=5),
        'caption': ParagraphStyle('caption',fontName='TNR-Italic',fontSize=10,leading=11,spaceAfter=5),
    }
    story=[]
    doc=Document()
    section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=Mm(13);section.bottom_margin=Mm(13)
    section.left_margin=Mm(15);section.right_margin=Mm(15)
    for style in ('Normal','Title','Heading 1','Caption'):
        s=doc.styles[style]
        s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(14)
    for kind,text in blocks:
        if kind=='break':
            story.append(PageBreak()); doc.add_page_break(); continue
        if kind=='image':
            image=ROOT/'screenshots'/text
            width,height=PILImage.open(image).size
            story.append(Image(str(image),width=180*72/25.4,height=180*72/25.4*height/width))
            paragraph=doc.add_paragraph()
            paragraph.paragraph_format.space_after=Pt(1)
            paragraph.paragraph_format.space_before=Pt(0)
            paragraph.paragraph_format.line_spacing=1
            paragraph.add_run().add_picture(str(image),width=Mm(180))
            continue
        story.append(Paragraph(escape(text),styles[kind]))
        paragraph=doc.add_paragraph(style={'title':'Title','heading':'Heading 1','caption':'Caption'}.get(kind,'Normal'))
        pf=paragraph.paragraph_format
        pf.space_before=Pt(5 if kind=='heading' else 0)
        pf.space_after=Pt(5 if kind in ('body','heading','caption') else 2)
        pf.line_spacing=Pt({'body':16,'heading':16,'header':14,'title':16,'caption':11}[kind])
        if kind in ('header','title'): paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=paragraph.add_run(text)
        run.font.name='Times New Roman';run.font.size=Pt({'body':14,'heading':14,'header':12,'title':14,'caption':10}[kind])
        run.bold=kind in ('heading','title');run.italic=kind=='caption'
    # Remove template paragraph borders, including the default Title rule.
    for element in (doc.styles.element, doc.element):
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.author='Sofiya Serbina'
    doc.core_properties.title='Laboratory Work No. 2 - REST API and Message Queues'
    doc.save(args.docx)
    SimpleDocTemplate(str(ROOT/'report.pdf'),pagesize=(210*72/25.4,297*72/25.4),
        leftMargin=15*72/25.4,rightMargin=15*72/25.4,topMargin=13*72/25.4,bottomMargin=13*72/25.4,
        title='Laboratory Work No. 2 - REST API and Message Queues',author='Sofiya Serbina').build(story)
    print('Created', ROOT/'report.pdf', 'and', args.docx)


if __name__=='__main__':
    main()
